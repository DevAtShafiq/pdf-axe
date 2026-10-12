"""Build the preview-panel test files into a folder (argv[1]).

    python qa/e2e/make_preview_fixtures.py <out_dir>

Creates: multi-page PDF with English + Korean text and a web link, a locked
PDF, a broken PDF, an empty PDF-named file, images (png/jpg/tiff/webp/gif),
UTF-8 and CP949 text, a .docx with an http and a mailto link, a multi-sheet
.xlsx, a .csv in Korean, an unknown extension, and a PowerPoint deck (via
PowerPoint itself when it is installed).
"""
from __future__ import annotations

import os
import sys


def main(out: str) -> None:
    import pymupdf as fitz
    from PIL import Image
    os.makedirs(out, exist_ok=True)

    # PDF: 3 pages, Korean text, a URI link
    doc = fitz.open()
    for i in range(3):
        p = doc.new_page()
        p.insert_text((72, 90), f"Page {i + 1} Hello world", fontsize=14)
        p.insert_text((72, 130), "안녕하세요 한국어 텍스트", fontsize=14, fontname="korea")
        if i == 0:
            p.insert_text((72, 170), "Visit https://example.com/docs now", fontsize=12)
            r = fitz.Rect(72, 190, 200, 206)
            p.insert_text((72, 202), "Click here", fontsize=12)
            p.insert_link({"kind": fitz.LINK_URI, "from": r, "uri": "https://example.org/x"})
    doc.save(os.path.join(out, "sample.pdf"))
    doc.close()

    d = fitz.open()
    d.new_page().insert_text((72, 72), "secret", fontsize=12)
    d.save(os.path.join(out, "locked.pdf"), encryption=fitz.PDF_ENCRYPT_AES_256,
           owner_pw="own", user_pw="pw")
    d.close()
    with open(os.path.join(out, "broken.pdf"), "wb") as f:
        f.write(b"%PDF-1.4\nthis is not really a pdf\n")
    open(os.path.join(out, "empty.pdf"), "wb").close()

    # images
    im = Image.new("RGB", (800, 500), (30, 120, 200))
    im.save(os.path.join(out, "img.png"))
    im.save(os.path.join(out, "img.jpg"))
    im.save(os.path.join(out, "img.tiff"))
    im.save(os.path.join(out, "img.tif"), format="TIFF")
    im.save(os.path.join(out, "img.webp"))
    im.convert("P").save(os.path.join(out, "img.gif"))
    with open(os.path.join(out, "broken.png"), "wb") as f:
        f.write(b"\x89PNG\r\n\x1a\nnope")

    # text
    with open(os.path.join(out, "utf8.txt"), "w", encoding="utf-8") as f:
        f.write("Hello\n안녕하세요 세계\n")
    with open(os.path.join(out, "korean_cp949.txt"), "wb") as f:
        f.write("한글 CP949 텍스트\n".encode("cp949"))

    # docx
    import docx
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    dd = docx.Document()
    dd.add_heading("제목 Heading", 1)
    dd.add_paragraph("Body text 본문.")

    def link(par, url, text):
        rid = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                                 is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), rid)
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        r.append(t)
        h.append(r)
        par._p.append(h)
    link(dd.add_paragraph("Web: "), "https://example.com/word", "web link")
    link(dd.add_paragraph("Mail: "), "mailto:someone@example.com", "mail link")
    t = dd.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "A1"
    t.cell(1, 1).text = "B2"
    dd.save(os.path.join(out, "doc.docx"))
    with open(os.path.join(out, "broken.docx"), "wb") as f:
        f.write(b"PK\x03\x04garbage")

    # xlsx
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["이름", "Amount", "Date"])
    import datetime
    ws.append(["홍길동", 1234.5, datetime.date(2026, 1, 2)])
    ws["B2"].number_format = "#,##0.00"
    ws2 = wb.create_sheet("Second")
    ws2["A1"] = "on sheet two"
    ws3 = wb.create_sheet("Hidden")
    ws3.sheet_state = "hidden"
    wb.save(os.path.join(out, "book.xlsx"))
    with open(os.path.join(out, "data.csv"), "w", encoding="utf-8-sig") as f:
        f.write("이름,점수\n김철수,90\n이영희,85\n")
    with open(os.path.join(out, "file.xyz"), "wb") as f:
        f.write(b"unknown")

    # pptx via PowerPoint (optional)
    if "--no-office" not in sys.argv:
        try:
            import pythoncom
            import win32com.client
            pythoncom.CoInitialize()
            app = win32com.client.DispatchEx("PowerPoint.Application")
            pres = app.Presentations.Add(WithWindow=False)
            s = pres.Slides.Add(1, 1)
            s.Shapes(1).TextFrame.TextRange.Text = "슬라이드 Slide One"
            pres.SaveAs(os.path.abspath(os.path.join(out, "deck.pptx")))
            pres.Close()
            app.Quit()
        except Exception as exc:   # PowerPoint missing
            print("pptx skipped:", exc)
    print("fixtures in", out)


if __name__ == "__main__":
    main(sys.argv[1])
