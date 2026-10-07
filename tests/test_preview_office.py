"""Preview pane: selectable PDF text layer, Word / Excel / CSV previews and the
Microsoft Office → PDF fallback (doc_preview.py + doc_preview_bridge.py)."""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import os

import pytest

pymupdf = pytest.importorskip("pymupdf")
docx = pytest.importorskip("docx")
openpyxl = pytest.importorskip("openpyxl")
from PIL import Image  # noqa: E402

import doc_preview as dp  # noqa: E402
from doc_preview_bridge import DocPreviewBridgeMixin  # noqa: E402


def _sha(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


# ─────────────────────────────────────────────────────────────────────────────
# PDF text layer
# ─────────────────────────────────────────────────────────────────────────────

@pytest.fixture()
def text_pdf(tmp_path):
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((72, 100), "Hello world from Office Axe", fontsize=14)
    p.insert_text((72, 140), "Passport number A1234567", fontsize=11)
    p2 = doc.new_page()                       # rotated page
    p2.insert_text((72, 100), "Rotated page text", fontsize=20)
    p2.set_rotation(90)
    p3 = doc.new_page()                       # scanned: image only
    buf = io.BytesIO()
    Image.effect_noise((300, 200), 40).convert("RGB").save(buf, "PNG")
    p3.insert_image(pymupdf.Rect(50, 50, 500, 350), stream=buf.getvalue())
    doc.new_page()                            # blank page
    out = tmp_path / "text.pdf"
    doc.save(str(out))
    return str(out)


def _words(layer):
    return [w for line in layer["lines"] for w in line["w"]]


def _search_rects(path, page, needle):
    with pymupdf.open(path) as d:
        pg = d[page]
        return [r * pg.rotation_matrix for r in pg.search_for(needle)]


def test_text_layer_words_match_search_boxes(text_pdf):
    layer = dp.pdf_text_layer(text_pdf, 0)
    assert layer["has_text"] and not layer["scanned"]
    assert (layer["width"], layer["height"]) == (595.0, 842.0)
    words = _words(layer)
    assert [w[0] for w in words] == ["Hello", "world", "from", "Office", "Axe",
                                     "Passport", "number", "A1234567"]
    # Space flags: first word of a line has none, the rest do.
    assert [w[5] for w in words[:5]] == [0, 1, 1, 1, 1]
    for text, x, y, length, h, _ in words:
        (r,) = _search_rects(text_pdf, 0, text)
        assert abs(x - r.x0) < 1.0, text
        assert abs((x + length) - r.x1) < 1.5, text
        assert abs(y - r.y0) < 1.0, text
        assert abs((y + h) - r.y1) < 1.5, text
    # Two lines → two line entries (copy keeps the line break)
    assert len(layer["lines"]) == 2 and all(l["a"] == 0 for l in layer["lines"])


def test_text_layer_rotated_page(text_pdf):
    layer = dp.pdf_text_layer(text_pdf, 1)
    assert layer["rotation"] == 90
    assert (layer["width"], layer["height"]) == (842.0, 595.0)   # rotated page size
    line = layer["lines"][0]
    assert line["a"] == 90.0
    words = line["w"]
    assert [w[0] for w in words] == ["Rotated", "page", "text"]
    for text, x, y, length, h, _ in words:
        (r,) = _search_rects(text_pdf, 1, text)
        # rotate(90deg) about (x, y): the word runs down from y and its height
        # extends to the left of x — i.e. the box [x-h, x] × [y, y+length].
        assert abs(x - r.x1) < 1.0, text
        assert abs((x - h) - r.x0) < 1.5, text
        assert abs(y - r.y0) < 1.0, text
        assert abs((y + length) - r.y1) < 1.5, text


def test_text_layer_image_only_page_is_flagged_scanned(text_pdf):
    layer = dp.pdf_text_layer(text_pdf, 2)
    assert layer["lines"] == [] and layer["words"] == 0
    assert layer["has_text"] is False and layer["scanned"] is True
    blank = dp.pdf_text_layer(text_pdf, 3)
    assert blank["has_text"] is False and blank["scanned"] is False


def test_text_layer_cache_follows_file_changes(text_pdf, tmp_path):
    first = dp.pdf_text_layer(text_pdf, 0)
    assert dp.pdf_text_layer(text_pdf, 0) is first           # cached
    import file_ops as fo
    fo.invalidate_pdf_doc_cache(text_pdf)
    with pymupdf.open(text_pdf) as d:
        d[0].set_rotation(180)
        d.save(str(tmp_path / "rot.pdf"))
    os.replace(str(tmp_path / "rot.pdf"), text_pdf)          # same path, new content
    st = os.stat(text_pdf)
    os.utime(text_pdf, ns=(st.st_atime_ns, st.st_mtime_ns + 5_000_000_000))
    fo.invalidate_pdf_doc_cache(text_pdf)
    again = dp.pdf_text_layer(text_pdf, 0)
    assert again is not first and again["rotation"] == 180


def test_page_and_all_text(text_pdf):
    t = dp.pdf_page_text(text_pdf, 0)
    assert "Hello world from Office Axe" in t and "Passport number A1234567" in t
    assert t.index("Hello") < t.index("Passport") and "\n" in t
    all_text, pages = dp.pdf_all_text(text_pdf)
    assert pages == 4 and "Rotated page text" in all_text and all_text == all_text.strip()


def test_text_layer_bridge_errors(text_pdf, tmp_path):
    b = DocPreviewBridgeMixin()
    r = b.pdf_text_layer(text_pdf, 99)
    assert r["ok"] is False and r["code"] == "range"
    r = b.pdf_text_layer(str(tmp_path / "missing.pdf"), 0)
    assert r["ok"] is False and r["code"] == "not_found"
    ok = b.pdf_page_text(text_pdf, 0)
    assert ok["ok"] and ok["chars"] == len(ok["text"]) > 10
    allr = b.pdf_all_text(text_pdf)
    assert allr["ok"] and allr["pages"] == 4


def test_text_layer_locked_pdf(tmp_path):
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "secret text", fontsize=12)
    p = tmp_path / "locked.pdf"
    doc.save(str(p), encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="pw2")
    r = DocPreviewBridgeMixin().pdf_text_layer(str(p), 0)
    assert r["ok"] is False and r["code"] == "locked"
    # After the preview unlock (in memory, file unchanged) the text layer works.
    import sfm_bridge
    b = sfm_bridge.SFMBridge()
    assert b.pdf_unlock_preview(str(p), "pw")["ok"]
    # Re-selecting the file asks for the state again; that must not re-lock /
    # corrupt the unlocked document (PyMuPDF's needs_pass resets the key).
    st = b.pdf_preview_state(str(p))
    assert st["ok"] and st["locked"] is False and st["needs_password"] is True
    assert b.pdf_unlock_preview(str(p), "pw")["ok"]          # unlocking twice is harmless
    r = b.pdf_text_layer(str(p), 0)
    assert r["ok"] and [w[0] for w in r["lines"][0]["w"]] == ["secret", "text"]
    import file_ops as fo
    fo.invalidate_pdf_doc_cache(str(p))


# ─────────────────────────────────────────────────────────────────────────────
# Word
# ─────────────────────────────────────────────────────────────────────────────

@pytest.fixture()
def sample_docx(tmp_path):
    from docx.shared import Inches
    d = docx.Document()
    d.add_heading("Quarterly Report", 1)
    d.add_heading("Summary section", 2)
    p = d.add_paragraph("Plain then ")
    p.add_run("bold words").bold = True
    p.add_run(" and ")
    p.add_run("italic part").italic = True
    p.add_run(" and ")
    p.add_run("underlined").underline = True
    p.add_run(" <script>x</script>")                       # must be escaped
    d.add_paragraph("first bullet", style="List Bullet")
    d.add_paragraph("second bullet", style="List Bullet")
    d.add_paragraph("step one", style="List Number")
    t = d.add_table(rows=3, cols=3)
    t.cell(0, 0).merge(t.cell(0, 1)).text = "Merged header"
    t.cell(1, 0).text = "Cell R1C1"
    t.cell(1, 1).merge(t.cell(2, 1)).text = "Tall cell"
    buf = io.BytesIO()
    Image.new("RGB", (40, 30), (200, 30, 30)).save(buf, "PNG")
    buf.seek(0)
    d.add_picture(buf, width=Inches(1))
    out = tmp_path / "report.docx"
    d.save(str(out))
    return str(out)


def test_docx_html_has_headings_runs_lists_table_image(sample_docx):
    before = _sha(sample_docx)
    r = dp.docx_to_html(sample_docx)
    h = r["html"]
    assert "<h1>Quarterly Report</h1>" in h
    assert "<h2>Summary section</h2>" in h
    assert "<strong>bold words</strong>" in h and "<em>italic part</em>" in h and "<u>underlined</u>" in h
    assert "<script>" not in h and "&lt;script&gt;" in h
    assert "<ul><li>first bullet</li><li>second bullet</li></ul>" in h
    assert '<ol class="doc-ol-decimal"><li>step one</li></ol>' in h
    assert '<td colspan="2"><p>Merged header</p></td>' in h
    assert '<td rowspan="2"><p>Tall cell</p></td>' in h
    assert h.count("<tr>") == 3
    assert '<img class="doc-img"' in h and "data:image/png;base64," in h
    s = r["stats"]
    assert s["tables"] == 1 and s["images"] == 1 and s["words"] >= 15
    assert r["truncated"] is False
    assert _sha(sample_docx) == before                         # read-only


def test_docx_long_document_is_capped(tmp_path):
    d = docx.Document()
    for i in range(60):
        d.add_paragraph(f"Paragraph {i}")
    p = tmp_path / "long.docx"
    d.save(str(p))
    r = dp.docx_to_html(str(p), max_blocks=10)
    assert r["truncated"] is True
    assert "Paragraph 9" in r["html"] and "Paragraph 10" not in r["html"]


def test_docx_bridge_and_bad_file(sample_docx, tmp_path):
    b = DocPreviewBridgeMixin()
    r = b.docx_preview_html(sample_docx)
    assert r["ok"] and "<h1>" in r["html"] and r["app"] == "Word"
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a zip")
    r = b.docx_preview_html(str(bad))
    assert r["ok"] is False and r["code"] == "bad_file"


# ─────────────────────────────────────────────────────────────────────────────
# Excel / CSV
# ─────────────────────────────────────────────────────────────────────────────

@pytest.fixture()
def sample_xlsx(tmp_path):
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = "Item"; ws["A1"].font = Font(bold=True)
    ws["B1"] = "Amount"; ws["C1"] = "Date"; ws["D1"] = "Share"
    ws["A2"] = "Rent"; ws["B2"] = 1234.5; ws["B2"].number_format = "#,##0.00"
    ws["C2"] = dt.datetime(2026, 3, 4); ws["C2"].number_format = "yyyy-mm-dd"
    ws["D2"] = 0.256; ws["D2"].number_format = "0.0%"
    ws["A3"] = "Total row"; ws.merge_cells("A3:C3")
    ws["A3"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["D3"] = 42
    ws2 = wb.create_sheet("Notes")
    ws2["A1"] = "second sheet"
    out = tmp_path / "budget.xlsx"
    wb.save(str(out))
    return str(out)


def test_xlsx_sheets_and_formatted_cells(sample_xlsx):
    before = _sha(sample_xlsx)
    r = dp.xlsx_to_html(sample_xlsx)
    assert [s["name"] for s in r["sheets"]] == ["Budget", "Notes"]
    assert r["sheet"] == "Budget" and r["rows"] == 3 and r["cols"] == 4
    h = r["html"]
    assert '<th class="xl-ch">A</th>' in h and '<th class="xl-ch">D</th>' in h
    assert '<th class="xl-rh">3</th>' in h
    assert '<td style="font-weight:700">Item</td>' in h
    assert ">1,234.50<" in h and ">2026-03-04<" in h and ">25.6%<" in h and ">42<" in h
    assert '<td colspan="3" style="background:#FFFF00">Total row</td>' in h
    assert r["truncated"] is False
    r2 = dp.xlsx_to_html(sample_xlsx, "Notes")
    assert r2["sheet"] == "Notes" and ">second sheet<" in r2["html"]
    assert _sha(sample_xlsx) == before


def test_xlsx_cap_rows_and_cols(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 31):
        for c in range(1, 13):
            ws.cell(r, c, r * c)
    p = tmp_path / "big.xlsx"
    wb.save(str(p))
    r = dp.xlsx_to_html(str(p), max_rows=10, max_cols=5)
    assert r["truncated"] is True and r["rows_shown"] == 10 and r["cols_shown"] == 5
    assert r["html"].count("<tr>") == 11                      # header + 10 rows
    assert '<th class="xl-ch">F</th>' not in r["html"]


@pytest.mark.parametrize("value,fmt,expected", [
    (1234.5, "#,##0.00", "1,234.50"),
    (1234.5, "0", "1235"),
    (0.5, "0%", "50%"),
    (-5, '"$"#,##0.00_);[Red]("$"#,##0.00)', "($5.00)"),
    (12345678, "0.00E+00", "1.23E+07"),
    (3, "General", "3"),
    (2.5, "General", "2.5"),
    (True, "General", "TRUE"),
    (None, "General", ""),
    (dt.datetime(2026, 3, 4, 14, 5), "m/d/yyyy h:mm", "3/4/2026 14:05"),
    (dt.datetime(2026, 3, 4), "d-mmm-yy", "4-Mar-26"),
    (dt.time(13, 30), "h:mm AM/PM", "1:30 PM"),
    (dt.datetime(2026, 3, 4), "General", "2026-03-04"),
])
def test_format_cell_value(value, fmt, expected):
    assert dp.format_cell_value(value, fmt) == expected


def test_csv_grid(tmp_path):
    p = tmp_path / "people.csv"
    p.write_text("name;age\nAlice;30\nBob;41\n", encoding="utf-8")
    r = DocPreviewBridgeMixin().sheet_preview_html(str(p))
    assert r["ok"] and r["delimiter"] == ";" and r["rows"] == 3 and r["cols"] == 2
    assert "<td>Alice</td>" in r["html"] and '<td class="xl-num">41</td>' in r["html"]


def test_quick_stats_and_file_info(sample_xlsx, sample_docx, text_pdf):
    assert dp.quick_stats(sample_xlsx) == {"sheets": ["Budget", "Notes"]}
    import sfm_bridge
    b = sfm_bridge.SFMBridge()
    info = b.get_file_info(text_pdf)
    assert info["ok"] and info["pages"] == 4
    info = b.get_file_info(sample_xlsx)
    assert info["sheets"] == ["Budget", "Notes"]


# ─────────────────────────────────────────────────────────────────────────────
# Office → PDF fallback (fake converter — no Office needed)
# ─────────────────────────────────────────────────────────────────────────────

def _fake_converter(calls):
    def conv(app, src, dst):
        calls.append((app, src, dst))
        d = pymupdf.open()
        d.new_page().insert_text((72, 72), f"Converted by {app}", fontsize=12)
        d.save(dst)
    return conv


def test_office_to_pdf_caches_and_never_touches_source(tmp_path, monkeypatch):
    monkeypatch.setenv("OFFICEAXE_PREVIEW_CACHE", str(tmp_path / "cache"))
    src = tmp_path / "old.doc"
    src.write_bytes(b"\xd0\xcf\x11\xe0 fake word 97 file")
    before = _sha(src)
    calls = []
    r = dp.office_to_pdf(str(src), _converter=_fake_converter(calls))
    assert r["app"] == "Word" and r["pages"] == 1 and r["cached"] is False
    assert os.path.dirname(r["pdf_path"]) == str(tmp_path / "cache")
    r2 = dp.office_to_pdf(str(src), _converter=_fake_converter(calls))
    assert r2["cached"] is True and r2["pdf_path"] == r["pdf_path"] and len(calls) == 1
    assert _sha(src) == before
    # The converted PDF gets a text layer like any other PDF
    assert "Converted" in dp.pdf_page_text(r["pdf_path"], 0)
    # A changed source → a new conversion
    src.write_bytes(b"\xd0\xcf\x11\xe0 changed content!!")
    st = os.stat(src)
    os.utime(src, ns=(st.st_atime_ns, st.st_mtime_ns + 5_000_000_000))
    r3 = dp.office_to_pdf(str(src), _converter=_fake_converter(calls))
    assert r3["cached"] is False and r3["pdf_path"] != r["pdf_path"] and len(calls) == 2


def test_office_to_pdf_timeout_and_failure(tmp_path, monkeypatch):
    import threading
    monkeypatch.setenv("OFFICEAXE_PREVIEW_CACHE", str(tmp_path / "cache"))
    src = tmp_path / "deck.pptx"
    src.write_bytes(b"PK fake pptx")
    gate = threading.Event()

    def slow(app, s, d):
        gate.wait(5)

    with pytest.raises(dp.PreviewError) as ei:
        dp.office_to_pdf(str(src), timeout=0.3, _converter=slow)
    assert ei.value.code == "timeout" and "PowerPoint" in str(ei.value)
    gate.set()

    def broken(app, s, d):
        raise RuntimeError("boom")

    src2 = tmp_path / "sheet.xls"
    src2.write_bytes(b"fake xls")
    with pytest.raises(dp.PreviewError) as ei:
        dp.office_to_pdf(str(src2), _converter=broken)
    assert ei.value.code == "convert_failed" and "Excel" in str(ei.value)


def test_office_app_mapping_and_info(tmp_path):
    assert dp.office_app_for(".doc") == "Word"
    assert dp.office_app_for(".xls") == "Excel"
    assert dp.office_app_for(".pptx") == "PowerPoint"
    assert dp.office_app_for(".docx") is None
    assert dp.office_app_for(".docx", include_native=True) == "Word"
    assert dp.office_app_for(".zip") is None
    p = tmp_path / "x.zip"
    p.write_bytes(b"PK")
    r = DocPreviewBridgeMixin().office_preview_info(str(p))
    assert r["ok"] and r["app"] == "" and r["available"] is False
