"""Regression tests for preview-panel bugs found in the QA pass (qa/preview.md)."""
from __future__ import annotations

import os
import sys

import pytest

pymupdf = pytest.importorskip("pymupdf")

import doc_preview as dp  # noqa: E402


@pytest.mark.skipif(sys.platform != "win32", reason="Windows path rules")
def test_office_gets_normalised_windows_paths(tmp_path, monkeypatch):
    # PowerPoint's SaveAs fails ("can't save ^0 to ^1") on forward slashes.
    monkeypatch.setenv("OFFICEAXE_PREVIEW_CACHE", str(tmp_path / "cache").replace("\\", "/"))
    src = tmp_path / "deck.pptx"
    src.write_bytes(b"PK fake pptx")
    seen = []

    def conv(app, s, d):
        seen.append((s, d))
        doc = pymupdf.open()
        doc.new_page()
        doc.save(d)

    r = dp.office_to_pdf(str(src).replace("\\", "/"), _converter=conv)
    assert r["pages"] == 1
    s, d = seen[0]
    assert "/" not in s and "/" not in d
    assert "/" not in dp.cached_pdf_path(str(src))


def test_broken_image_preview_has_a_readable_error(tmp_path):
    from sfm_bridge import SFMBridge
    p = tmp_path / "broken.png"
    p.write_bytes(b"\x89PNG\r\n\x1a\nnope")
    r = SFMBridge.get_image_preview(object.__new__(SFMBridge), str(p), 400)
    assert r["ok"] is False and "NoneType" not in r["error"] and "damaged" in r["error"]
    r = SFMBridge.get_image_preview(object.__new__(SFMBridge), str(tmp_path / "gone.png"), 400)
    assert r["ok"] is False and "not found" in r["error"]


def test_broken_docx_error_does_not_dump_the_path(tmp_path):
    pytest.importorskip("docx")          # without python-docx the error is 'missing_dep'
    p = tmp_path / "broken.docx"
    p.write_bytes(b"PK\x03\x04garbage")
    with pytest.raises(dp.PreviewError) as ei:
        dp.docx_to_html(str(p))
    assert ei.value.code == "bad_file"
    assert str(tmp_path) not in str(ei.value) and "damaged" in str(ei.value)


def _big_wb(tmp_path, monkeypatch):
    openpyxl = pytest.importorskip("openpyxl")
    p = tmp_path / "big.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "top"
    ws["A5"] = "five"
    ws["C8"] = 3
    wb.save(str(p))
    real = os.path.getsize
    monkeypatch.setattr(dp.os.path, "getsize",
                        lambda q: 99 * 1024 * 1024 if os.path.abspath(q) == os.path.abspath(str(p)) else real(q))
    return str(p)


def test_big_workbook_row_numbers_are_right(tmp_path, monkeypatch):
    import re
    p = _big_wb(tmp_path, monkeypatch)       # > 15 MB → read-only mode
    r = dp.xlsx_to_html(p)
    assert re.findall(r'<th class="xl-rh">(\d+)</th>', r["html"]) == [str(i) for i in range(1, 9)]
    assert ">five<" in r["html"]


def test_big_workbook_preview_does_not_lock_the_file(tmp_path, monkeypatch):
    p = _big_wb(tmp_path, monkeypatch)
    dp.xlsx_to_html(p)
    os.rename(p, p + ".moved")               # WinError 32 while the preview held it
    os.rename(p + ".moved", p)


def test_big_workbook_without_dimension_is_capped(tmp_path, monkeypatch):
    import re
    import zipfile
    openpyxl = pytest.importorskip("openpyxl")
    a, b = tmp_path / "a.xlsx", tmp_path / "nodim.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(1, 31):
        ws.cell(i, 1, f"r{i}")
        ws.cell(i, 2, i)
    wb.save(str(a))
    with zipfile.ZipFile(a) as zin, zipfile.ZipFile(b, "w") as zout:
        for it in zin.infolist():
            d = zin.read(it.filename)
            if it.filename.startswith("xl/worksheets/"):
                d = re.sub(rb"<dimension[^>]*/>", b"", d)
            zout.writestr(it, d)
    real = os.path.getsize
    monkeypatch.setattr(dp.os.path, "getsize",
                        lambda q: 99 * 1024 * 1024 if os.path.abspath(q) == os.path.abspath(str(b)) else real(q))
    r = dp.xlsx_to_html(str(b), max_rows=10, max_cols=5)
    assert r["rows_shown"] == 10 and r["cols_shown"] == 2
    assert r["rows"] >= 10 and r["cols"] == 2


def test_previewed_pdf_is_not_locked_and_reloads_after_external_save(tmp_path):
    from sfm_bridge import SFMBridge
    b = SFMBridge.__new__(SFMBridge)
    p = str(tmp_path / "ext.pdf")
    d = pymupdf.open()
    d.new_page().insert_text((72, 72), "version one")
    d.save(p)
    d.close()
    assert SFMBridge.get_pdf_page_as_png(b, p, 0, 72)["ok"]
    assert dp.pdf_page_text(p, 0) == "version one"
    # Another program saves over the file while it is in the preview.
    d = pymupdf.open()
    for i in range(2):
        d.new_page().insert_text((72, 72), f"version two p{i}")
    d.save(p)                                   # was "Permission denied"
    d.close()
    assert dp.pdf_page_text(p, 0) == "version two p0"
    assert SFMBridge.get_pdf_page_as_png(b, p, 1, 72)["ok"]
    os.rename(p, p + ".x")
    os.rename(p + ".x", p)


def test_docx_only_web_links_are_clickable(tmp_path):
    docx = pytest.importorskip("docx")
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d = docx.Document()
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
    for url, text in (("https://example.com/word", "web link"), ("mailto:a@example.com", "mail link")):
        par = d.add_paragraph()
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), par.part.relate_to(url, rel, is_external=True))
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        r.append(t)
        h.append(r)
        par._p.append(h)
    p = tmp_path / "links.docx"
    d.save(str(p))
    html = dp.docx_to_html(str(p))["html"]
    assert 'data-href="https://example.com/word"' in html
    assert "mailto:" not in html and "mail link" in html
