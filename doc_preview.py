"""
doc_preview.py — read-only preview helpers for the preview pane.

  * PDF text layer: positioned words for a transparent, selectable layer over
    the rendered page image (like pdf.js), plus plain page / document text.
  * Word (.docx) → HTML via python-docx (headings, runs, lists, tables with
    merged cells, inline images as data URLs).
  * Excel (.xlsx/.xlsm) and .csv → HTML grid (sheet tabs, column letters,
    formatted values, bold/fill colours, merged cells).
  * Other Office files (.doc/.xls/.pptx …) → PDF through Microsoft Office
    (win32com) into %LOCALAPPDATA%\\OfficeAxe\\preview_cache, shown in the
    normal PDF preview.

Nothing here ever modifies the source file.
"""
from __future__ import annotations

import base64
import csv
import datetime as _dt
import hashlib
import html
import io
import math
import os
import re
import threading
import time
import zipfile
from collections import OrderedDict

import file_ops as _fo

# ═════════════════════════════════════════════════════════════════════════════
# PDF TEXT LAYER
# ═════════════════════════════════════════════════════════════════════════════

_TL_CACHE: "OrderedDict[tuple, dict]" = OrderedDict()
_TL_CACHE_MAX = 96
_TL_LOCK = threading.Lock()


class PreviewError(Exception):
    def __init__(self, msg: str, code: str = ""):
        super().__init__(msg)
        self.code = code


def _file_sig(path: str) -> tuple:
    st = os.stat(path)
    return (os.path.abspath(path).lower(), st.st_mtime_ns, st.st_size)


def _open_pdf(path: str):
    """The preview's cached (possibly password-unlocked) document."""
    if not os.path.isfile(path):
        raise PreviewError(f"File not found: {os.path.basename(path)}", "not_found")
    doc = _fo.get_cached_fitz_doc(path)
    if doc is None:
        raise PreviewError("Could not open the PDF", "open_failed")
    if getattr(doc, "needs_pass", False) and getattr(doc, "is_encrypted", False):
        raise PreviewError("The PDF is locked", "locked")
    return doc


def _r(v: float) -> float:
    return round(float(v), 2)


def _page_words(page) -> list[dict]:
    """Lines of words in display (rotated-page) coordinates, in points.

    Each line: {"a": angle_deg, "w": [[text, x, y, length, height, space_before], …]}
    (x, y) is the top-left corner of the word box in its own (rotated) frame,
    so CSS ``left/top`` + ``transform: rotate(a) scaleX(..)`` with
    ``transform-origin: 0 0`` puts the word exactly over the rendered glyphs.
    """
    import pymupdf as fitz  # noqa: F401  (installed as pymupdf/fitz)
    rm = page.rotation_matrix
    flags = 0
    for name in ("TEXT_PRESERVE_WHITESPACE", "TEXT_PRESERVE_LIGATURES", "TEXT_MEDIABOX_CLIP"):
        flags |= int(getattr(fitz, name, 0))
    raw = page.get_text("rawdict", flags=flags)
    out: list[dict] = []
    for block in raw.get("blocks", []):
        if block.get("type", 0) != 0:
            continue
        for line in block.get("lines", []):
            dx, dy = line.get("dir", (1.0, 0.0))
            ddx = dx * rm.a + dy * rm.c
            ddy = dx * rm.b + dy * rm.d
            n = math.hypot(ddx, ddy) or 1.0
            ddx, ddy = ddx / n, ddy / n
            upx, upy = ddy, -ddx            # "up" in y-down display space
            angle = math.degrees(math.atan2(ddy, ddx))

            def proj(px, py):
                return px * ddx + py * ddy

            words: list[list] = []
            cur = None
            space = False

            def flush():
                nonlocal cur
                if cur is not None:
                    ox, oy = cur["o"]
                    a, h = cur["asc"], cur["h"]
                    x = ox + upx * a
                    y = oy + upy * a
                    length = max(cur["e"] - cur["s"], 0.5)
                    words.append([cur["t"], _r(x), _r(y), _r(length), _r(h), 1 if cur["sp"] else 0])
                    cur = None

            for span in line.get("spans", []):
                size = float(span.get("size") or 0) or 1.0
                asc = float(span.get("ascender", 0.9) or 0.9)
                desc = float(span.get("descender", -0.25) or -0.25)
                if not (0.5 <= asc <= 1.3):
                    asc = 0.9
                if not (-0.6 <= desc <= 0.0):
                    desc = -0.25
                for ch in span.get("chars", []):
                    c = ch.get("c", "")
                    if not c or c.isspace():
                        flush()
                        space = True
                        continue
                    o = fitz.Point(ch["origin"]) * rm
                    bb = fitz.Rect(ch["bbox"]) * rm
                    s = proj(o.x, o.y)
                    e = max(proj(bb.x0, bb.y0), proj(bb.x1, bb.y0), proj(bb.x0, bb.y1), proj(bb.x1, bb.y1))
                    if e <= s:
                        e = s + size * 0.3
                    # A big gap without a space character still separates words.
                    if cur is not None and s - cur["e"] > size * 0.3:
                        flush()
                        space = True
                    if cur is None:
                        cur = {"t": c, "o": (o.x, o.y), "s": s, "e": e,
                               "asc": asc * size, "h": (asc - desc) * size,
                               "sp": space and bool(words)}
                        space = False
                    else:
                        cur["t"] += c
                        cur["e"] = max(cur["e"], e)
                        cur["asc"] = max(cur["asc"], asc * size)
                        cur["h"] = max(cur["h"], (asc - desc) * size)
            flush()
            if words:
                out.append({"a": _r(angle), "w": words})
    return out


def pdf_text_layer(path: str, page_index: int) -> dict:
    """Positioned words for page *page_index* (0-based), cached per file version.

    → {width, height, rotation, lines, words, has_text, scanned}
    """
    page_index = int(page_index)
    key = _file_sig(path) + (page_index,)
    with _TL_LOCK:
        hit = _TL_CACHE.get(key)
        if hit is not None:
            _TL_CACHE.move_to_end(key)
            return hit
    doc = _open_pdf(path)
    if page_index < 0 or page_index >= len(doc):
        raise PreviewError(f"Page {page_index + 1} is out of range", "range")
    page = doc.load_page(page_index)
    rect = page.rect
    lines = _page_words(page)
    nwords = sum(len(l["w"]) for l in lines)
    scanned = False
    if not nwords:
        try:
            scanned = bool(page.get_images(full=False)) or bool(page.get_image_info())
        except Exception:
            scanned = False
    res = {"width": _r(rect.width), "height": _r(rect.height), "rotation": int(page.rotation),
           "lines": lines, "words": nwords, "has_text": nwords > 0, "scanned": scanned}
    with _TL_LOCK:
        _TL_CACHE[key] = res
        while len(_TL_CACHE) > _TL_CACHE_MAX:
            _TL_CACHE.popitem(last=False)
    return res


def pdf_page_text(path: str, page_index: int) -> str:
    doc = _open_pdf(path)
    page_index = int(page_index)
    if page_index < 0 or page_index >= len(doc):
        raise PreviewError(f"Page {page_index + 1} is out of range", "range")
    return _clean_text(doc.load_page(page_index).get_text("text", sort=True))


def pdf_all_text(path: str, max_chars: int = 20_000_000) -> tuple[str, int]:
    doc = _open_pdf(path)
    parts: list[str] = []
    total = 0
    for i in range(len(doc)):
        t = _clean_text(doc.load_page(i).get_text("text", sort=True))
        parts.append(t)
        total += len(t)
        if total > max_chars:
            break
    return "\n\n".join(parts).rstrip(), len(doc)


def _clean_text(t: str) -> str:
    t = (t or "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [ln.rstrip() for ln in t.split("\n")]
    return "\n".join(lines).strip("\n")


# ═════════════════════════════════════════════════════════════════════════════
# WORD (.docx) → HTML
# ═════════════════════════════════════════════════════════════════════════════

DOCX_EXT = {".docx", ".dotx"}
SHEET_EXT = {".xlsx", ".xlsm", ".xltx", ".xltm"}
CSV_EXT = {".csv", ".tsv"}

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
_V = "{urn:schemas-microsoft-com:vml}"

MAX_IMAGE_BYTES = 3 * 1024 * 1024        # per image
MAX_IMAGES_TOTAL = 12 * 1024 * 1024      # per document
MAX_DOCX_BLOCKS = 4000                   # paragraphs + tables at top level
MAX_DOCX_HTML = 6 * 1024 * 1024

_IMG_MIME = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif",
             ".bmp": "image/bmp", ".webp": "image/webp", ".svg": "image/svg+xml"}


def _e(s) -> str:
    return html.escape(str(s or ""), quote=True)


def _wattr(el, name):
    return el.get(_W + name) if el is not None else None


class _DocxRenderer:
    def __init__(self, document):
        self.doc = document
        self.part = document.part
        self.img_total = 0
        self.images = 0
        self.images_skipped = 0
        self.tables = 0
        self.words = 0
        self.paragraphs = 0
        self._num_cache: dict = {}
        self._style_cache: dict = {}

    # ── numbering (lists) ────────────────────────────────────────────────────
    def _num_fmt(self, num_id: str, ilvl: str) -> str:
        key = (num_id, ilvl)
        if key in self._num_cache:
            return self._num_cache[key]
        fmt = "bullet"
        try:
            numbering = self.part.numbering_part.element
            abs_id = None
            for num in numbering.findall(_W + "num"):
                if num.get(_W + "numId") == num_id:
                    a = num.find(_W + "abstractNumId")
                    abs_id = _wattr(a, "val")
                    break
            if abs_id is not None:
                for an in numbering.findall(_W + "abstractNum"):
                    if an.get(_W + "abstractNumId") == abs_id:
                        for lvl in an.findall(_W + "lvl"):
                            if lvl.get(_W + "ilvl") == ilvl:
                                nf = lvl.find(_W + "numFmt")
                                fmt = _wattr(nf, "val") or "bullet"
                                break
                        break
        except Exception:
            pass
        self._num_cache[key] = fmt
        return fmt

    def _para_num(self, p_el):
        """(numId, ilvl) for a list paragraph — direct or through its style."""
        ppr = p_el.find(_W + "pPr")
        numpr = ppr.find(_W + "numPr") if ppr is not None else None
        if numpr is None and ppr is not None:
            ps = ppr.find(_W + "pStyle")
            sid = _wattr(ps, "val")
            if sid:
                numpr = self._style_numpr(sid)
        if numpr is None:
            return None
        nid = _wattr(numpr.find(_W + "numId"), "val")
        lvl = _wattr(numpr.find(_W + "ilvl"), "val") or "0"
        if not nid or nid == "0":
            return None
        return nid, lvl

    def _style_numpr(self, style_id):
        if style_id in self._style_cache:
            return self._style_cache[style_id]
        res = None
        try:
            st = self.doc.styles.element
            for s in st.findall(_W + "style"):
                if s.get(_W + "styleId") == style_id:
                    ppr = s.find(_W + "pPr")
                    if ppr is not None:
                        res = ppr.find(_W + "numPr")
                    break
        except Exception:
            res = None
        self._style_cache[style_id] = res
        return res

    # ── blocks ───────────────────────────────────────────────────────────────
    def render_blocks(self, parent_el, out: list, top_level=False, limit=None) -> bool:
        """Render paragraphs/tables of *parent_el*. Returns True when truncated."""
        list_stack: list[tuple[str, int]] = []   # (tag, level)
        count = 0

        def close_lists(to_level=-1):
            while list_stack and list_stack[-1][1] > to_level:
                tag, _ = list_stack.pop()
                out.append(f"</li></{tag}>")

        children = list(parent_el.iterchildren())
        for child in children:
            tag = child.tag
            if tag == _W + "sdt":
                content = child.find(_W + "sdtContent")
                if content is not None:
                    self.render_blocks(content, out)
                continue
            if tag == _W + "p":
                count += 1
                num = self._para_num(child)
                if num:
                    nid, lvl = num
                    level = int(lvl) if str(lvl).isdigit() else 0
                    # "List Bullet 2" / "List Number 3" styles nest by name
                    sm = re.match(r"list (?:bullet|number|continue|paragraph)\s*(\d)$",
                                  self._para_style(child)[0].lower())
                    if sm and level == 0:
                        level = max(0, int(sm.group(1)) - 1)
                    fmt = self._num_fmt(nid, lvl)
                    ltag = "ul" if fmt in ("bullet", "none") else "ol"
                    cls = f' class="doc-ol-{_e(fmt)}"' if ltag == "ol" else ""
                    if list_stack and list_stack[-1][1] > level:
                        close_lists(level)
                    if list_stack and list_stack[-1][1] == level:
                        if list_stack[-1][0] == ltag:
                            out.append("</li>")
                        else:
                            close_lists(level - 1)
                    if not list_stack or list_stack[-1][1] < level:
                        out.append(f"<{ltag}{cls}>")
                        list_stack.append((ltag, level))
                    out.append("<li>")
                    out.append(self.render_paragraph(child, inner_only=True))
                else:
                    close_lists()
                    out.append(self.render_paragraph(child))
            elif tag == _W + "tbl":
                count += 1
                close_lists()
                out.append(self.render_table(child))
            else:
                continue
            if top_level and limit and count >= limit:
                close_lists()
                return True
            if top_level and count % 100 == 0 and sum(len(x) for x in out) > MAX_DOCX_HTML:
                close_lists()
                return True
        close_lists()
        return False

    def _para_style(self, p_el):
        ppr = p_el.find(_W + "pPr")
        sid = _wattr(ppr.find(_W + "pStyle"), "val") if ppr is not None else None
        name = ""
        if sid:
            try:
                st = self.doc.styles.element
                for s in st.findall(_W + "style"):
                    if s.get(_W + "styleId") == sid:
                        nm = s.find(_W + "name")
                        name = _wattr(nm, "val") or sid
                        break
            except Exception:
                name = sid
            name = name or sid
        return (name or "").strip(), ppr

    def render_paragraph(self, p_el, inner_only=False) -> str:
        self.paragraphs += 1
        name, ppr = self._para_style(p_el)
        lname = name.lower()
        tag = "p"
        m = re.match(r"heading\s*(\d)", lname)
        if m:
            tag = "h" + str(max(1, min(6, int(m.group(1)))))
        elif lname == "title":
            tag = "h1"
        elif lname == "subtitle":
            tag = "h2"
        style = []
        if ppr is not None:
            jc = _wattr(ppr.find(_W + "jc"), "val")
            if jc in ("center", "right", "both", "end", "start", "left"):
                style.append("text-align:" + {"both": "justify", "end": "right", "start": "left"}.get(jc, jc))
            ind = ppr.find(_W + "ind")
            if ind is not None and not inner_only:
                left = _wattr(ind, "left") or _wattr(ind, "start")
                try:
                    if left and int(left) > 0:
                        style.append(f"margin-left:{int(left) / 20:.0f}pt")
                except ValueError:
                    pass
            shd = ppr.find(_W + "shd")
            fill = _wattr(shd, "fill")
            if fill and re.fullmatch(r"[0-9A-Fa-f]{6}", fill):
                style.append(f"background:#{fill}")
        inner = self.render_runs(p_el)
        text = re.sub(r"<[^>]+>", "", inner)
        self.words += len(html.unescape(text).split())
        if inner_only:
            return inner or "&#8203;"
        cls = ""
        if lname in ("quote", "intense quote"):
            cls = ' class="doc-quote"'
        st = f' style="{";".join(style)}"' if style else ""
        if not inner.strip():
            return f"<{tag}{cls}{st}>&nbsp;</{tag}>" if tag == "p" else ""
        return f"<{tag}{cls}{st}>{inner}</{tag}>"

    def render_runs(self, parent, href=None) -> str:
        parts: list[str] = []
        for el in parent.iterchildren():
            t = el.tag
            if t == _W + "r":
                parts.append(self.render_run(el))
            elif t == _W + "hyperlink":
                link = None
                rid = el.get(_R + "id")
                if rid:
                    try:
                        link = self.part.rels[rid].target_ref
                    except Exception:
                        link = None
                inner = self.render_runs(el)
                if link and re.match(r"^(https?:|mailto:)", link, re.I):
                    parts.append(f'<a data-href="{_e(link)}" class="doc-link" title="{_e(link)}">{inner}</a>')
                else:
                    parts.append(inner)
            elif t in (_W + "ins", _W + "smartTag", _W + "fldSimple", _W + "customXml"):
                parts.append(self.render_runs(el))
            elif t == _W + "sdt":
                c = el.find(_W + "sdtContent")
                if c is not None:
                    parts.append(self.render_runs(c))
        return "".join(parts)

    def render_run(self, r_el) -> str:
        rpr = r_el.find(_W + "rPr")
        b = i = u = s = False
        va = None
        color = None
        hl = None
        size = None
        if rpr is not None:
            def on(name):
                x = rpr.find(_W + name)
                if x is None:
                    return False
                v = _wattr(x, "val")
                return v not in ("0", "false", "off", "none")
            b, i, s = on("b"), on("i"), on("strike") or on("dstrike")
            ue = rpr.find(_W + "u")
            u = ue is not None and _wattr(ue, "val") != "none"
            va = _wattr(rpr.find(_W + "vertAlign"), "val")
            c = _wattr(rpr.find(_W + "color"), "val")
            if c and re.fullmatch(r"[0-9A-Fa-f]{6}", c) and c.upper() != "000000":
                color = c
            h = _wattr(rpr.find(_W + "highlight"), "val")
            if h and h != "none":
                hl = h
            sz = _wattr(rpr.find(_W + "sz"), "val")
            try:
                size = int(sz) / 2 if sz else None
            except ValueError:
                size = None
        segs: list[str] = []
        for el in r_el.iterchildren():
            t = el.tag
            if t == _W + "t":
                segs.append(_e(el.text or ""))
            elif t == _W + "tab":
                segs.append('<span class="doc-tab">\t</span>')
            elif t in (_W + "br", _W + "cr"):
                if _wattr(el, "type") == "page":
                    segs.append('<span class="doc-pagebreak"></span>')
                else:
                    segs.append("<br>")
            elif t == _W + "noBreakHyphen":
                segs.append("&#8209;")
            elif t == _W + "sym":
                segs.append("&#9632;")
            elif t in (_W + "drawing", _W + "pict"):
                segs.append(self.render_image(el))
            elif t == "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent":
                segs.append(self.render_image(el))
        txt = "".join(segs)
        if not txt:
            return ""
        if b:
            txt = f"<strong>{txt}</strong>"
        if i:
            txt = f"<em>{txt}</em>"
        if u:
            txt = f"<u>{txt}</u>"
        if s:
            txt = f"<s>{txt}</s>"
        if va == "superscript":
            txt = f"<sup>{txt}</sup>"
        elif va == "subscript":
            txt = f"<sub>{txt}</sub>"
        st = []
        if color:
            st.append(f"color:#{color}")
        if hl:
            st.append(f"background:{_e(hl)}")
        if size and abs(size - 11) > 0.6 and 5 <= size <= 72:
            st.append(f"font-size:{size:g}pt")
        if st:
            txt = f'<span style="{";".join(st)}">{txt}</span>'
        return txt

    def render_image(self, el) -> str:
        blips = el.findall(".//" + _A + "blip")
        out = []
        for blip in blips:
            rid = blip.get(_R + "embed")
            if not rid:
                continue
            w = h = None
            ext_el = el.find(".//" + _WP + "extent")
            if ext_el is not None:
                try:
                    w = int(ext_el.get("cx")) / 9525
                    h = int(ext_el.get("cy")) / 9525
                except (TypeError, ValueError):
                    w = h = None
            try:
                ip = self.part.related_parts[rid]
                blob = ip.blob
                ext = os.path.splitext(str(getattr(ip, "partname", "")))[1].lower()
            except Exception:
                continue
            mime = _IMG_MIME.get(ext)
            if not mime:
                blob, mime = self._convert_image(blob)
                if not blob:
                    self.images_skipped += 1
                    out.append('<span class="doc-img-missing">[image]</span>')
                    continue
            if len(blob) > MAX_IMAGE_BYTES or self.img_total + len(blob) > MAX_IMAGES_TOTAL:
                small = self._shrink(blob)
                if small is None or self.img_total + len(small) > MAX_IMAGES_TOTAL:
                    self.images_skipped += 1
                    out.append('<span class="doc-img-missing">[image too large to preview]</span>')
                    continue
                blob, mime = small, "image/jpeg"
            self.img_total += len(blob)
            self.images += 1
            dim = ""
            if w and h:
                dim = f' style="width:{w:.0f}px;max-width:100%;height:auto;aspect-ratio:{w:.0f}/{h:.0f}"'
            out.append(f'<img class="doc-img" alt=""{dim} src="data:{mime};base64,'
                       f'{base64.b64encode(blob).decode("ascii")}">')
            break   # AlternateContent may repeat the same picture
        if not out:
            # VML picture (old .doc-style images inside .docx)
            for im in el.findall(".//" + _V + "imagedata"):
                rid = im.get(_R + "id")
                if rid:
                    try:
                        blob = self.part.related_parts[rid].blob
                        small = self._shrink(blob)
                        if small and self.img_total + len(small) <= MAX_IMAGES_TOTAL:
                            self.img_total += len(small)
                            self.images += 1
                            out.append('<img class="doc-img" alt="" src="data:image/jpeg;base64,'
                                       + base64.b64encode(small).decode("ascii") + '">')
                            break
                    except Exception:
                        pass
        return "".join(out)

    @staticmethod
    def _shrink(blob: bytes, max_dim: int = 1400):
        try:
            from PIL import Image
            im = Image.open(io.BytesIO(blob))
            im.load()
            if im.mode not in ("RGB", "L"):
                bg = Image.new("RGB", im.size, (255, 255, 255))
                try:
                    bg.paste(im, mask=im.convert("RGBA").split()[-1])
                except Exception:
                    bg.paste(im.convert("RGB"))
                im = bg
            im.thumbnail((max_dim, max_dim))
            buf = io.BytesIO()
            im.save(buf, "JPEG", quality=82)
            return buf.getvalue()
        except Exception:
            return None

    def _convert_image(self, blob: bytes):
        small = self._shrink(blob)
        return (small, "image/jpeg") if small else (None, None)

    def render_table(self, tbl_el) -> str:
        self.tables += 1
        rows = tbl_el.findall(_W + "tr")
        # grid[r] = list of (col_start, span, tc_el, vmerge)
        layout = []
        for tr in rows:
            col = 0
            cells = []
            for tc in tr.findall(_W + "tc"):
                tcpr = tc.find(_W + "tcPr")
                span = 1
                vm = None
                if tcpr is not None:
                    gs = _wattr(tcpr.find(_W + "gridSpan"), "val")
                    try:
                        span = max(1, int(gs)) if gs else 1
                    except ValueError:
                        span = 1
                    vme = tcpr.find(_W + "vMerge")
                    if vme is not None:
                        vm = _wattr(vme, "val") or "continue"
                cells.append([col, span, tc, vm, 1])
                col += span
            layout.append(cells)
        # resolve vertical merges → rowspan on the restarting cell
        for ri, cells in enumerate(layout):
            for cell in cells:
                if cell[3] == "continue":
                    for up in range(ri - 1, -1, -1):
                        origin = next((c for c in layout[up] if c[0] == cell[0]), None)
                        if origin is None:
                            break
                        if origin[3] != "continue":
                            origin[4] += 1
                            break
        out = ['<table class="doc-table"><tbody>']
        for cells in layout:
            out.append("<tr>")
            for col, span, tc, vm, rowspan in cells:
                if vm == "continue":
                    continue
                attrs = ""
                if span > 1:
                    attrs += f' colspan="{span}"'
                if rowspan > 1:
                    attrs += f' rowspan="{rowspan}"'
                tcpr = tc.find(_W + "tcPr")
                if tcpr is not None:
                    fill = _wattr(tcpr.find(_W + "shd"), "fill")
                    if fill and re.fullmatch(r"[0-9A-Fa-f]{6}", fill) and fill.upper() != "FFFFFF":
                        attrs += f' style="background:#{fill}"'
                inner: list[str] = []
                self.render_blocks(tc, inner)
                out.append(f"<td{attrs}>{''.join(inner)}</td>")
            out.append("</tr>")
        out.append("</tbody></table>")
        return "".join(out)


def docx_app_props(path: str) -> dict:
    """Pages / Words / Slides … from docProps/app.xml (Office's last save)."""
    res: dict = {}
    try:
        with zipfile.ZipFile(path) as z:
            data = z.read("docProps/app.xml").decode("utf-8", "replace")
        for k in ("Pages", "Words", "Characters", "Paragraphs", "Slides", "Lines"):
            m = re.search(rf"<{k}>(\d+)</{k}>", data)
            if m:
                res[k.lower()] = int(m.group(1))
    except Exception:
        pass
    return res


def docx_to_html(path: str, max_blocks: int = MAX_DOCX_BLOCKS) -> dict:
    """→ {html, truncated, stats{paragraphs, words, tables, images, pages?}}"""
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise PreviewError("python-docx is not installed", "missing_dep") from exc
    try:
        document = docx.Document(path)
    except Exception as exc:
        raise PreviewError(f"Could not read the Word document: {exc}", "bad_file") from exc
    rnd = _DocxRenderer(document)
    out: list[str] = []
    body = document.element.body
    truncated = rnd.render_blocks(body, out, top_level=True, limit=max_blocks)
    # Remaining top-level blocks after truncation, for the notice
    total_blocks = sum(1 for c in body.iterchildren() if c.tag in (_W + "p", _W + "tbl"))
    stats = {"paragraphs": rnd.paragraphs, "words": rnd.words, "tables": rnd.tables,
             "images": rnd.images, "images_skipped": rnd.images_skipped, "blocks": total_blocks}
    props = docx_app_props(path)
    if "pages" in props:
        stats["pages"] = props["pages"]
    if truncated:
        # Count the words of the full document cheaply for the stats.
        try:
            stats["words"] = sum(len((p.text or "").split()) for p in document.paragraphs)
        except Exception:
            pass
    return {"html": "".join(out), "truncated": bool(truncated), "stats": stats}


# ═════════════════════════════════════════════════════════════════════════════
# EXCEL (.xlsx/.xlsm) / CSV → HTML grid
# ═════════════════════════════════════════════════════════════════════════════

MAX_ROWS = 1000
MAX_COLS = 60
_WB_CACHE: "OrderedDict[tuple, object]" = OrderedDict()
_WB_LOCK = threading.Lock()


def col_letter(n: int) -> str:
    """1 → A, 27 → AA."""
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _load_wb(path: str):
    key = _file_sig(path)
    with _WB_LOCK:
        wb = _WB_CACHE.get(key)
        if wb is not None:
            _WB_CACHE.move_to_end(key)
            return wb
    import openpyxl
    big = os.path.getsize(path) > 15 * 1024 * 1024
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=big, keep_links=False)
    except Exception as exc:
        raise PreviewError(f"Could not read the workbook: {exc}", "bad_file") from exc
    with _WB_LOCK:
        _WB_CACHE[key] = wb
        while len(_WB_CACHE) > 2:
            _, old = _WB_CACHE.popitem(last=False)
            try:
                old.close()
            except Exception:
                pass
    return wb


# ── number formats ──────────────────────────────────────────────────────────

_DATE_TOKEN_RE = re.compile(
    r'(\[[^\]]*\]|"[^"]*"|\\.|yyyy|yy|mmmmm|mmmm|mmm|mm|m|dddd|ddd|dd|d|hh|h|ss|s|AM/PM|am/pm|A/P|a/p|\.0+|.)')


def _strip_fmt_section(fmt: str, value) -> str:
    secs = _split_sections(fmt)
    if not secs:
        return "General"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if value < 0 and len(secs) > 1:
            return secs[1]
        if value == 0 and len(secs) > 2:
            return secs[2]
    return secs[0]


def _split_sections(fmt: str) -> list[str]:
    out, cur, q = [], "", False
    i = 0
    while i < len(fmt):
        c = fmt[i]
        if c == '"':
            q = not q
        if c == "\\" and i + 1 < len(fmt):
            cur += fmt[i:i + 2]
            i += 2
            continue
        if c == ";" and not q:
            out.append(cur)
            cur = ""
        else:
            cur += c
        i += 1
    out.append(cur)
    return out


def _format_date(value, fmt: str) -> str:
    if isinstance(value, _dt.time):
        value = _dt.datetime.combine(_dt.date(1900, 1, 1), value)
    elif isinstance(value, _dt.date) and not isinstance(value, _dt.datetime):
        value = _dt.datetime.combine(value, _dt.time())
    f = re.sub(r"\[(?!h\]|m\]|s\])[^\]]*\]", "", fmt)      # colours / locales
    tokens = _DATE_TOKEN_RE.findall(f)
    ampm = any(t.lower() in ("am/pm", "a/p") for t in tokens)
    out = []
    last_h = False
    for idx, t in enumerate(tokens):
        tl = t.lower()
        if t.startswith('"'):
            out.append(t[1:-1])
        elif t.startswith("\\"):
            out.append(t[1:])
        elif tl == "yyyy":
            out.append(f"{value.year:04d}")
        elif tl == "yy":
            out.append(f"{value.year % 100:02d}")
        elif tl in ("mm", "m"):
            # minutes after an hour token or before a seconds token
            nxt = next((x.lower() for x in tokens[idx + 1:] if x.strip() and x not in (":", ".")), "")
            if last_h or nxt.startswith("s"):
                out.append(f"{value.minute:02d}" if tl == "mm" else str(value.minute))
            else:
                out.append(f"{value.month:02d}" if tl == "mm" else str(value.month))
        elif tl == "mmmmm":
            out.append(value.strftime("%B")[0])
        elif tl == "mmmm":
            out.append(value.strftime("%B"))
        elif tl == "mmm":
            out.append(value.strftime("%b"))
        elif tl == "dddd":
            out.append(value.strftime("%A"))
        elif tl == "ddd":
            out.append(value.strftime("%a"))
        elif tl == "dd":
            out.append(f"{value.day:02d}")
        elif tl == "d":
            out.append(str(value.day))
        elif tl in ("hh", "h"):
            hr = value.hour
            if ampm:
                hr = hr % 12 or 12
            out.append(f"{hr:02d}" if tl == "hh" else str(hr))
            last_h = True
            continue
        elif tl in ("ss", "s"):
            out.append(f"{value.second:02d}" if tl == "ss" else str(value.second))
        elif tl in ("am/pm", "a/p"):
            pm = value.hour >= 12
            out.append(("PM" if pm else "AM") if tl == "am/pm" else ("P" if pm else "A"))
        elif t.startswith(".0"):
            out.append("." + f"{value.microsecond:06d}"[:len(t) - 1])
        elif t in ("[h]",):
            out.append(str(value.hour))
        else:
            out.append(t)
        if t.strip() and t not in (":",):
            last_h = False if tl not in ("hh", "h") else last_h
    return "".join(out)


def _format_number(value: float, fmt: str) -> str:
    sec = _strip_fmt_section(fmt, value)
    neg_section = value < 0 and len(_split_sections(fmt)) > 1
    v = abs(value) if neg_section else value
    s = re.sub(r"\[[^\]]*\]", lambda m: m.group(0)[2:-1].split("-")[0] if m.group(0).startswith("[$") else "", sec)
    # literal text
    pre, post, core = "", "", s
    m = re.search(r"[0#?][0#?,.]*(?:E[+-]0+)?%?", s, re.I)
    if not m:
        lit = re.sub(r'"([^"]*)"', r"\1", s).replace("\\", "").replace("_)", "").replace("*", "")
        return lit.strip() or _general(value)
    pre = s[:m.start()]
    core = m.group(0)
    post = s[m.end():]

    def lit(t):
        t = re.sub(r'"([^"]*)"', r"\1", t)
        t = re.sub(r"_.", " ", t)
        t = re.sub(r"\*.", "", t)
        return t.replace("\\", "")
    pct = "%" in core or "%" in post
    if pct:
        v = v * 100
    if re.search(r"E[+-]", core, re.I):
        dec = len(core.split(".")[1].split("E")[0].split("e")[0]) if "." in core else 0
        body = f"{v:.{dec}E}"
        mant, exp = body.split("E")
        body = f"{mant}E{exp[0]}{int(exp[1:]):02d}"
    else:
        num = core.rstrip("%")
        dec_part = num.split(".")[1] if "." in num else ""
        dec = len(re.sub(r"[^0#?]", "", dec_part))
        min_dec = len(re.sub(r"[^0]", "", dec_part))
        int_part = num.split(".")[0]
        # trailing commas scale by 1000
        while int_part.endswith(","):
            v = v / 1000
            int_part = int_part[:-1]
        grouping = "," in int_part
        body = f"{v:,.{dec}f}" if grouping else f"{v:.{dec}f}"
        if dec > min_dec and "." in body:
            body = body.rstrip("0")
            if len(body.split(".")[1]) < min_dec:
                body += "0" * (min_dec - len(body.split(".")[1]))
            body = body.rstrip(".")
        min_int = len(re.sub(r"[^0]", "", int_part))
        if min_int == 0 and body.startswith("0.") :
            body = body[1:]
        elif min_int == 0 and body in ("0", "-0"):
            body = ""
    out = lit(pre) + body + ("%" if "%" in core else "") + lit(post)
    if neg_section and "(" not in out and value < 0 and "-" not in out:
        pass
    return out.strip()


def _general(value) -> str:
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return str(value)
        if value == int(value) and abs(value) < 1e15:
            return str(int(value))
        r = f"{value:.10g}"
        return r
    return str(value)


def format_cell_value(value, number_format: str = "General") -> str:
    """Display text for a cell value using its Excel number format (best effort)."""
    if value is None:
        return ""
    fmt = number_format or "General"
    try:
        if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
            if fmt == "General" or not re.search(r"[ymdhs]", re.sub(r'"[^"]*"', "", fmt), re.I):
                if isinstance(value, _dt.datetime):
                    return value.strftime("%Y-%m-%d %H:%M" if (value.hour or value.minute or value.second) else "%Y-%m-%d")
                if isinstance(value, _dt.time):
                    return value.strftime("%H:%M:%S")
                return value.strftime("%Y-%m-%d")
            return _format_date(value, _strip_fmt_section(fmt, 0))
        if isinstance(value, _dt.timedelta):
            secs = int(value.total_seconds())
            return f"{secs // 3600}:{secs % 3600 // 60:02d}:{secs % 60:02d}"
        if isinstance(value, bool):
            return "TRUE" if value else "FALSE"
        if isinstance(value, (int, float)):
            if fmt in ("General", "@", ""):
                return _general(value)
            return _format_number(float(value), fmt)
    except Exception:
        return _general(value)
    return str(value)


def _rgb(color) -> str | None:
    """'#RRGGBB' for an openpyxl Color with an explicit RGB (themes ignored)."""
    try:
        if color is None or getattr(color, "type", None) != "rgb":
            return None
        v = str(color.rgb or "")
        if len(v) == 8 and re.fullmatch(r"[0-9A-Fa-f]{8}", v):
            if v[:2] == "00" and v[2:].upper() == "000000":
                return None
            return "#" + v[2:]
        if len(v) == 6 and re.fullmatch(r"[0-9A-Fa-f]{6}", v):
            return "#" + v
    except Exception:
        pass
    return None


def _sheet_extent(ws, max_rows: int, max_cols: int) -> tuple[int, int, int, int]:
    """(rows_used, cols_used, rows_shown, cols_shown) — trailing blanks trimmed."""
    mr = ws.max_row or 0
    mc = ws.max_column or 0
    last_r = last_c = 0
    for row in ws.iter_rows(min_row=1, max_row=min(mr, max_rows), max_col=min(mc, max_cols)):
        for c in row:
            if c.value is not None and c.value != "":
                rr = getattr(c, "row", None)
                cc = getattr(c, "column", None)
                if rr and rr > last_r:
                    last_r = rr
                if cc and cc > last_c:
                    last_c = cc
    try:
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= max_rows and rng.min_col <= max_cols:
                last_r = max(last_r, min(rng.max_row, max_rows))
                last_c = max(last_c, min(rng.max_col, max_cols))
    except Exception:
        pass
    return mr, mc, last_r, last_c


def xlsx_to_html(path: str, sheet: str = "", max_rows: int = MAX_ROWS, max_cols: int = MAX_COLS) -> dict:
    """→ {html, sheets:[{name, hidden}], sheet, rows, cols, rows_shown, cols_shown, truncated}"""
    wb = _load_wb(path)
    sheets = []
    for ws in wb.worksheets:
        sheets.append({"name": ws.title, "hidden": getattr(ws, "sheet_state", "visible") != "visible"})
    if not sheets:
        raise PreviewError("This workbook has no worksheets", "empty")
    names = [s["name"] for s in sheets]
    if sheet not in names:
        try:
            active = wb.active.title if wb.active is not None else None
        except Exception:
            active = None
        sheet = active if active in names else next((s["name"] for s in sheets if not s["hidden"]), names[0])
    ws = wb[sheet]
    with _WB_LOCK:
        mr, mc, nr, nc = _sheet_extent(ws, max_rows, max_cols)
        html_out = _grid_html_ws(ws, nr, nc)
    truncated = mr > max_rows or mc > max_cols
    if truncated:
        # rows_used may be inflated by formatting; report the cap honestly
        pass
    return {"html": html_out, "sheets": sheets, "sheet": sheet, "rows": mr, "cols": mc,
            "rows_shown": nr, "cols_shown": nc,
            "truncated": bool(truncated and (mr > nr or mc > nc)),
            "max_rows": max_rows, "max_cols": max_cols}


def _grid_html_ws(ws, nr: int, nc: int) -> str:
    merged_origin: dict = {}
    merged_skip: set = set()
    try:
        for rng in ws.merged_cells.ranges:
            r0, c0 = rng.min_row, rng.min_col
            if r0 > nr or c0 > nc:
                continue
            r1, c1 = min(rng.max_row, nr), min(rng.max_col, nc)
            merged_origin[(r0, c0)] = (r1 - r0 + 1, c1 - c0 + 1)
            for r in range(r0, r1 + 1):
                for c in range(c0, c1 + 1):
                    if (r, c) != (r0, c0):
                        merged_skip.add((r, c))
    except Exception:
        pass
    out = ['<table class="xl-grid"><colgroup><col class="xl-col-rh">']
    for c in range(1, nc + 1):
        w = None
        try:
            cd = ws.column_dimensions.get(col_letter(c)) if hasattr(ws.column_dimensions, "get") else None
            if cd is not None and cd.width:
                w = int(cd.width * 7 + 5)
            if cd is not None and getattr(cd, "hidden", False):
                w = 0
        except Exception:
            w = None
        out.append(f'<col style="width:{max(24, min(w, 600))}px">' if w else "<col>")
    out.append('</colgroup><thead><tr><th class="xl-corner"></th>')
    for c in range(1, nc + 1):
        out.append(f'<th class="xl-ch">{col_letter(c)}</th>')
    out.append("</tr></thead><tbody>")
    if nr and nc:
        for row in ws.iter_rows(min_row=1, max_row=nr, max_col=nc):
            cells = list(row)
            rnum = None
            for c in cells:
                rnum = getattr(c, "row", None)
                if rnum:
                    break
            rnum = rnum or (len(out))
            out.append(f'<tr><th class="xl-rh">{rnum}</th>')
            for ci, cell in enumerate(cells, start=1):
                if (rnum, ci) in merged_skip:
                    continue
                attrs = ""
                span = merged_origin.get((rnum, ci))
                if span:
                    if span[0] > 1:
                        attrs += f' rowspan="{span[0]}"'
                    if span[1] > 1:
                        attrs += f' colspan="{span[1]}"'
                v = getattr(cell, "value", None)
                cls = []
                st = []
                try:
                    nf = getattr(cell, "number_format", "General") or "General"
                    text = format_cell_value(v, nf)
                    font = getattr(cell, "font", None)
                    if font is not None:
                        if font.b:
                            st.append("font-weight:700")
                        if font.i:
                            st.append("font-style:italic")
                        if font.u and font.u != "none":
                            st.append("text-decoration:underline")
                        col = _rgb(font.color)
                        if col:
                            st.append(f"color:{col}")
                    fill = getattr(cell, "fill", None)
                    if fill is not None and getattr(fill, "fill_type", None) == "solid":
                        bg = _rgb(fill.fgColor)
                        if bg and bg.upper() != "#FFFFFF":
                            st.append(f"background:{bg}")
                    al = getattr(cell, "alignment", None)
                    h = getattr(al, "horizontal", None) if al is not None else None
                    if h in ("center", "centerContinuous"):
                        st.append("text-align:center")
                    elif h == "right":
                        st.append("text-align:right")
                    elif h == "left":
                        st.append("text-align:left")
                    elif isinstance(v, (int, float, _dt.date, _dt.time, _dt.timedelta)) and not isinstance(v, bool):
                        cls.append("xl-num")
                    if al is not None and getattr(al, "wrap_text", False):
                        cls.append("xl-wrap")
                except Exception:
                    text = _general(v) if v is not None else ""
                if cls:
                    attrs += f' class="{" ".join(cls)}"'
                if st:
                    attrs += f' style="{";".join(st)}"'
                out.append(f"<td{attrs}>{_e(text)}</td>")
            out.append("</tr>")
    out.append("</tbody></table>")
    return "".join(out)


def _read_text_any(path: str, limit: int = 8 * 1024 * 1024) -> str:
    with open(path, "rb") as f:
        raw = f.read(limit)
    for enc in ("utf-8-sig", "cp949", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def csv_to_html(path: str, max_rows: int = MAX_ROWS, max_cols: int = MAX_COLS) -> dict:
    text = _read_text_any(path)
    sample = text[:20000]
    delim = "\t" if path.lower().endswith(".tsv") else ","
    try:
        delim = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except Exception:
        pass
    rows = []
    total = 0
    maxc = 0
    for row in csv.reader(io.StringIO(text), delimiter=delim):
        total += 1
        maxc = max(maxc, len(row))
        if len(rows) < max_rows:
            rows.append(row[:max_cols])
    nc = min(maxc, max_cols)
    out = ['<table class="xl-grid"><colgroup><col class="xl-col-rh">' + "<col>" * nc +
           '</colgroup><thead><tr><th class="xl-corner"></th>']
    for c in range(1, nc + 1):
        out.append(f'<th class="xl-ch">{col_letter(c)}</th>')
    out.append("</tr></thead><tbody>")
    num_re = re.compile(r"^-?[\d,]*\.?\d+%?$")
    for i, row in enumerate(rows, start=1):
        out.append(f'<tr><th class="xl-rh">{i}</th>')
        for c in range(nc):
            v = row[c] if c < len(row) else ""
            cls = ' class="xl-num"' if v and num_re.match(v.strip()) else ""
            out.append(f"<td{cls}>{_e(v)}</td>")
        out.append("</tr>")
    out.append("</tbody></table>")
    return {"html": "".join(out), "sheets": [{"name": os.path.basename(path), "hidden": False}],
            "sheet": os.path.basename(path), "rows": total, "cols": maxc,
            "rows_shown": len(rows), "cols_shown": nc,
            "truncated": total > max_rows or maxc > max_cols,
            "max_rows": max_rows, "max_cols": max_cols, "delimiter": delim}


def xlsx_sheet_names(path: str) -> list[str]:
    """Sheet names straight from workbook.xml (cheap — for the details pane)."""
    try:
        with zipfile.ZipFile(path) as z:
            data = z.read("xl/workbook.xml").decode("utf-8", "replace")
        return [html.unescape(n) for n in re.findall(r'<(?:\w+:)?sheet\b[^>]*\bname="([^"]*)"', data)]
    except Exception:
        return []


# ═════════════════════════════════════════════════════════════════════════════
# OTHER OFFICE FILES → PDF via Microsoft Office (win32com)
# ═════════════════════════════════════════════════════════════════════════════

WORD_EXT = {".doc", ".docm", ".dot", ".dotm", ".rtf", ".odt"}
EXCEL_EXT = {".xls", ".xlsb", ".xlt", ".ods"}
PPT_EXT = {".pptx", ".ppt", ".pptm", ".pps", ".ppsx", ".potx", ".odp"}
OFFICE_TIMEOUT = 120

_PROGIDS = {"Word": "Word.Application", "Excel": "Excel.Application", "PowerPoint": "PowerPoint.Application"}
_avail_cache: dict = {}
_convert_lock = threading.Lock()          # one Office automation at a time
_key_locks: dict = {}
_key_locks_lock = threading.Lock()


def office_app_for(ext: str, *, include_native: bool = False) -> str | None:
    ext = (ext or "").lower()
    if ext in WORD_EXT or (include_native and ext in DOCX_EXT):
        return "Word"
    if ext in EXCEL_EXT or (include_native and ext in (SHEET_EXT | {".csv"})):
        return "Excel"
    if ext in PPT_EXT:
        return "PowerPoint"
    return None


def office_available(app: str) -> bool:
    if app in _avail_cache:
        return _avail_cache[app]
    ok = False
    if os.name == "nt" and app in _PROGIDS:
        try:
            import winreg
            with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, _PROGIDS[app] + "\\CLSID"):
                ok = True
        except OSError:
            ok = False
        if ok:
            try:
                import win32com.client  # noqa: F401
            except Exception:
                ok = False
    _avail_cache[app] = ok
    return ok


def preview_cache_dir() -> str:
    base = os.environ.get("OFFICEAXE_PREVIEW_CACHE")
    if not base:
        root = os.environ.get("LOCALAPPDATA") or os.path.join(os.path.expanduser("~"), "AppData", "Local")
        base = os.path.join(root, "OfficeAxe", "preview_cache")
    os.makedirs(base, exist_ok=True)
    return base


def cached_pdf_path(path: str) -> str:
    sig = _file_sig(path)
    h = hashlib.sha1(f"{sig[0]}|{sig[1]}|{sig[2]}".encode("utf-8")).hexdigest()[:20]
    stem = re.sub(r"[^\w\-]+", "_", os.path.splitext(os.path.basename(path))[0])[:40] or "doc"
    return os.path.join(preview_cache_dir(), f"{stem}_{h}.pdf")


def _valid_pdf(p: str) -> int:
    try:
        if not os.path.isfile(p) or os.path.getsize(p) < 64:
            return 0
        import pymupdf as fitz
        with fitz.open(p) as d:
            return len(d)
    except Exception:
        return 0


def _convert_with_office(app: str, src: str, dst: str) -> None:
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    inst = None
    doc = None
    try:
        if app == "Word":
            inst = win32com.client.DispatchEx("Word.Application")
            inst.Visible = False
            try:
                inst.DisplayAlerts = 0
            except Exception:
                pass
            try:
                inst.AutomationSecurity = 3          # msoAutomationSecurityForceDisable (no macros)
            except Exception:
                pass
            doc = inst.Documents.Open(FileName=src, ConfirmConversions=False, ReadOnly=True,
                                      AddToRecentFiles=False, PasswordDocument="\u0001officeaxe",
                                      Revert=False, Visible=False, OpenAndRepair=False,
                                      NoEncodingDialog=True)
            doc.ExportAsFixedFormat(OutputFileName=dst, ExportFormat=17, OpenAfterExport=False,
                                    OptimizeFor=0, Range=0, IncludeDocProps=False,
                                    CreateBookmarks=0, DocStructureTags=True)
        elif app == "Excel":
            inst = win32com.client.DispatchEx("Excel.Application")
            inst.Visible = False
            inst.DisplayAlerts = False
            try:
                inst.AskToUpdateLinks = False
                inst.AutomationSecurity = 3
                inst.EnableEvents = False
            except Exception:
                pass
            doc = inst.Workbooks.Open(src, UpdateLinks=0, ReadOnly=True, Password="\u0001officeaxe",
                                      IgnoreReadOnlyRecommended=True, AddToMru=False, Notify=False)
            doc.ExportAsFixedFormat(0, dst)
        elif app == "PowerPoint":
            inst = win32com.client.DispatchEx("PowerPoint.Application")
            try:
                inst.DisplayAlerts = 1                # ppAlertsNone
                inst.AutomationSecurity = 3
            except Exception:
                pass
            doc = inst.Presentations.Open(src, ReadOnly=True, Untitled=False, WithWindow=False)
            doc.SaveAs(dst, 32)                       # ppSaveAsPDF
        else:
            raise PreviewError(f"No Office app for this file", "no_app")
    finally:
        if doc is not None:
            try:
                if app == "Word":
                    doc.Close(SaveChanges=0)
                elif app == "Excel":
                    doc.Close(SaveChanges=False)
                else:
                    doc.Close()
            except Exception:
                pass
        if inst is not None:
            try:
                inst.Quit()
            except Exception:
                pass
        doc = inst = None
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


def office_to_pdf(path: str, timeout: float = OFFICE_TIMEOUT, *, app: str | None = None,
                  _converter=None) -> dict:
    """Convert *path* to a cached PDF via Microsoft Office (background thread + timeout).

    → {pdf_path, pages, cached}. The source file is opened read-only and never changed.
    """
    if not os.path.isfile(path):
        raise PreviewError(f"File not found: {os.path.basename(path)}", "not_found")
    ext = os.path.splitext(path)[1].lower()
    app = app or office_app_for(ext, include_native=True)
    if not app:
        raise PreviewError("This file type has no Office preview", "unsupported")
    conv = _converter or _convert_with_office
    if _converter is None and not office_available(app):
        raise PreviewError(f"Microsoft {app} is not installed", "no_office")
    dst = cached_pdf_path(path)
    n = _valid_pdf(dst)
    if n:
        return {"pdf_path": dst, "pages": n, "cached": True, "app": app}
    with _key_locks_lock:
        lk = _key_locks.setdefault(dst, threading.Lock())
    with lk:
        n = _valid_pdf(dst)
        if n:
            return {"pdf_path": dst, "pages": n, "cached": True, "app": app}
        result: dict = {}

        def work():
            try:
                with _convert_lock:
                    conv(app, os.path.abspath(path), dst)
                result["ok"] = True
            except Exception as exc:     # noqa: BLE001
                result["error"] = exc

        t = threading.Thread(target=work, daemon=True, name="office-preview")
        t0 = time.time()
        t.start()
        t.join(timeout)
        if t.is_alive():
            raise PreviewError(f"Microsoft {app} took too long (over {int(timeout)} s)", "timeout")
        if "error" in result:
            exc = result["error"]
            msg = str(exc)
            if "password" in msg.lower():
                raise PreviewError("The document is password-protected", "password")
            raise PreviewError(f"Microsoft {app} could not convert the file: {msg[:300]}", "convert_failed")
        n = _valid_pdf(dst)
        if not n:
            raise PreviewError(f"Microsoft {app} did not produce a PDF", "convert_failed")
        return {"pdf_path": dst, "pages": n, "cached": False, "app": app,
                "seconds": round(time.time() - t0, 1)}


# ═════════════════════════════════════════════════════════════════════════════
# Details-pane counts (cheap)
# ═════════════════════════════════════════════════════════════════════════════

def quick_stats(path: str) -> dict:
    ext = os.path.splitext(path)[1].lower()
    res: dict = {}
    if ext in DOCX_EXT or ext in (".docm", ".pptx", ".pptm", ".ppsx"):
        props = docx_app_props(path)
        for k in ("pages", "words", "slides"):
            if props.get(k):
                res[k] = props[k]
    if ext in SHEET_EXT:
        names = xlsx_sheet_names(path)
        if names:
            res["sheets"] = names
    return res
